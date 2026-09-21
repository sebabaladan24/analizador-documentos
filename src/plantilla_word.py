"""Plantilla de Word "profesional" para los documentos PMI: portada con
título/proyecto/fecha, encabezados con color y línea separadora, pie de
página. Distinta a propósito del export genérico de la propuesta
(`exportar_word.markdown_a_docx`, texto plano convertido de markdown) —
acá sí importa que el resultado se vea como un documento de gestión de
proyectos real, no como texto corrido.
"""

import re
from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

COLOR_ACENTO = RGBColor(0x1F, 0x3A, 0x5F)
COLOR_SUBTITULO = RGBColor(0x55, 0x55, 0x55)
COLOR_FECHA = RGBColor(0x88, 0x88, 0x88)
COLOR_PIE = RGBColor(0xAA, 0xAA, 0xAA)


def _agregar_borde_inferior(paragraph, color: str = "1F3A5F", size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borde = OxmlElement("w:pBdr")
    inferior = OxmlElement("w:bottom")
    inferior.set(qn("w:val"), "single")
    inferior.set(qn("w:sz"), str(size))
    inferior.set(qn("w:space"), "4")
    inferior.set(qn("w:color"), color)
    borde.append(inferior)
    p_pr.append(borde)


def _estilo_base(documento: Document) -> None:
    estilo_normal = documento.styles["Normal"]
    estilo_normal.font.name = "Calibri"
    estilo_normal.font.size = Pt(11)


def _portada(documento: Document, titulo_documento: str, nombre_proyecto: str) -> None:
    documento.add_paragraph()

    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titulo = titulo.add_run(titulo_documento)
    run_titulo.font.size = Pt(26)
    run_titulo.font.bold = True
    run_titulo.font.color.rgb = COLOR_ACENTO

    subtitulo = documento.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_subtitulo = subtitulo.add_run(nombre_proyecto.strip() or "Proyecto sin nombre")
    run_subtitulo.font.size = Pt(16)
    run_subtitulo.font.color.rgb = COLOR_SUBTITULO

    fecha_parrafo = documento.add_paragraph()
    fecha_parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_fecha = fecha_parrafo.add_run(date.today().strftime("%d/%m/%Y"))
    run_fecha.font.size = Pt(11)
    run_fecha.font.color.rgb = COLOR_FECHA

    documento.add_paragraph()
    documento.add_page_break()


def _agregar_pie_pagina(documento: Document, titulo_documento: str) -> None:
    seccion = documento.sections[0]
    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pie.add_run(f"{titulo_documento} — generado el {date.today().strftime('%d/%m/%Y')}")
    run.font.size = Pt(8)
    run.font.color.rgb = COLOR_PIE


def _agregar_texto_con_negritas(parrafo, texto: str) -> None:
    partes = re.split(r"(\*\*.*?\*\*)", texto)
    for parte in partes:
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            parrafo.add_run(parte[2:-2]).bold = True
        elif parte:
            parrafo.add_run(parte)


def _agregar_cuerpo(documento: Document, markdown_texto: str) -> None:
    for linea in markdown_texto.split("\n"):
        linea = linea.rstrip()
        if not linea.strip():
            continue
        if linea.startswith("# "):
            continue  # el título ya está en la portada, no se repite
        if linea.startswith("## "):
            encabezado = documento.add_heading(linea[3:], level=1)
            for run in encabezado.runs:
                run.font.color.rgb = COLOR_ACENTO
            _agregar_borde_inferior(encabezado)
        elif linea.startswith("### "):
            documento.add_heading(linea[4:], level=2)
        elif linea.startswith("- ") or linea.startswith("* "):
            parrafo = documento.add_paragraph(style="List Bullet")
            _agregar_texto_con_negritas(parrafo, linea[2:])
        else:
            parrafo = documento.add_paragraph()
            parrafo.paragraph_format.space_after = Pt(6)
            _agregar_texto_con_negritas(parrafo, linea)


def documento_pmi_a_docx(titulo_documento: str, nombre_proyecto: str, contenido_markdown: str) -> BytesIO:
    documento = Document()
    _estilo_base(documento)
    _portada(documento, titulo_documento, nombre_proyecto)
    _agregar_cuerpo(documento, contenido_markdown)
    _agregar_pie_pagina(documento, titulo_documento)

    buffer = BytesIO()
    documento.save(buffer)
    buffer.seek(0)
    return buffer
