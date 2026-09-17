"""Conversión de la propuesta (markdown simple: #/##/###, listas con
'-', **negrita**) a un .docx editable con python-docx.
"""

import re
from io import BytesIO

from docx import Document


def _agregar_texto_con_negritas(parrafo, texto: str) -> None:
    partes = re.split(r"(\*\*.*?\*\*)", texto)
    for parte in partes:
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            parrafo.add_run(parte[2:-2]).bold = True
        elif parte:
            parrafo.add_run(parte)


def markdown_a_docx(markdown_texto: str) -> BytesIO:
    documento = Document()
    for linea in markdown_texto.split("\n"):
        linea = linea.rstrip()
        if not linea.strip():
            documento.add_paragraph("")
            continue
        if linea.startswith("### "):
            documento.add_heading(linea[4:], level=3)
        elif linea.startswith("## "):
            documento.add_heading(linea[3:], level=2)
        elif linea.startswith("# "):
            documento.add_heading(linea[2:], level=1)
        elif linea.startswith("- ") or linea.startswith("* "):
            parrafo = documento.add_paragraph(style="List Bullet")
            _agregar_texto_con_negritas(parrafo, linea[2:])
        else:
            parrafo = documento.add_paragraph()
            _agregar_texto_con_negritas(parrafo, linea)

    buffer = BytesIO()
    documento.save(buffer)
    buffer.seek(0)
    return buffer
