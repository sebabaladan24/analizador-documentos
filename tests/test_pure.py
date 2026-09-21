"""Tests de las partes puras del pipeline (sin Ollama ni Chroma en vivo)."""

from docx import Document

from src.documentos_pmi import PLANTILLAS
from src.exportar_word import markdown_a_docx
from src.extraccion import validar_requisitos
from src.ingesta import chunkear_texto
from src.plantilla_word import documento_pmi_a_docx
from src.servicios_extraccion import borrador_a_markdown, slug_desde_nombre
from src.servicios_loader import dividir_por_secciones


def test_chunkear_texto_vacio():
    assert chunkear_texto("") == []


def test_chunkear_texto_una_sola_parte_si_es_corto():
    texto = "una dos tres cuatro cinco"
    assert chunkear_texto(texto, tam_palabras=10, solape_palabras=2) == [texto]


def test_chunkear_texto_respeta_solape():
    palabras = [f"palabra{i}" for i in range(20)]
    texto = " ".join(palabras)
    chunks = chunkear_texto(texto, tam_palabras=10, solape_palabras=3)
    assert len(chunks) >= 2
    # el final del primer chunk se solapa con el inicio del segundo
    fin_chunk1 = chunks[0].split()[-3:]
    inicio_chunk2 = chunks[1].split()[:3]
    assert fin_chunk1 == inicio_chunk2


def test_dividir_por_secciones_basico():
    cuerpo = "## Descripción\ntexto uno\n\n## SLA\ntexto dos"
    secciones = dividir_por_secciones(cuerpo)
    assert secciones == [("Descripción", "texto uno"), ("SLA", "texto dos")]


def test_dividir_por_secciones_ignora_vacias():
    cuerpo = "\n\n## Sección A\ncontenido"
    secciones = dividir_por_secciones(cuerpo)
    assert secciones == [("Sección A", "contenido")]


def test_validar_requisitos_normaliza_categoria_invalida():
    entrada = [{"requisito": "Backup diario", "categoria": "no-existe"}]
    resultado = validar_requisitos(entrada)
    assert resultado == [{"requisito": "Backup diario", "categoria": "otro"}]


def test_validar_requisitos_ignora_items_malformados():
    entrada = [{"sin_requisito": True}, {"requisito": "Ancho de banda 500 Mbps", "categoria": "red"}]
    resultado = validar_requisitos(entrada)
    assert resultado == [{"requisito": "Ancho de banda 500 Mbps", "categoria": "red"}]


def test_validar_requisitos_rechaza_no_lista():
    try:
        validar_requisitos({"no": "es una lista"})
        assert False, "debería haber lanzado ValueError"
    except ValueError:
        pass


def test_plantillas_pmi_tiene_los_5_documentos_esperados():
    assert set(PLANTILLAS.keys()) == {
        "acta_inicio",
        "formulacion",
        "gestion_cambio",
        "lecciones_aprendidas",
        "cierre",
    }
    for plantilla in PLANTILLAS.values():
        assert plantilla["titulo"]
        assert len(plantilla["secciones"]) > 0


def test_slug_desde_nombre_normaliza():
    assert slug_desde_nombre("Nube Empresarial (IaaS)") == "nube-empresarial-iaas"
    assert slug_desde_nombre("   ") == "servicio"


def test_borrador_a_markdown_usa_placeholder_en_campos_vacios():
    datos = {
        "nombre": "Backup Extra",
        "categoria": "Continuidad",
        "descripcion": "Backup adicional para clientes grandes.",
        "especificaciones_tecnicas": "",
        "sla": "- Disponibilidad 99.9%",
        "cuando_ofrecerlo": "",
        "cuando_no_ofrecerlo": "",
    }
    markdown = borrador_a_markdown(datos)

    assert "nombre: Backup Extra" in markdown
    assert "categoria: Continuidad" in markdown
    assert "## Descripción\nBackup adicional para clientes grandes." in markdown
    assert "## Especificaciones técnicas\n[Verificar" in markdown
    assert "## SLA\n- Disponibilidad 99.9%" in markdown


def test_borrador_a_markdown_es_parseable_por_dividir_por_secciones():
    datos = {
        "nombre": "Servicio X",
        "categoria": "Cat",
        "descripcion": "Desc",
        "especificaciones_tecnicas": "Specs",
        "sla": "SLA",
        "cuando_ofrecerlo": "Ofrecer",
        "cuando_no_ofrecerlo": "No ofrecer",
    }
    markdown = borrador_a_markdown(datos)
    cuerpo = markdown.split("---\n", 2)[-1]
    secciones = dividir_por_secciones(cuerpo)
    titulos = [titulo for titulo, _ in secciones]
    assert titulos == [
        "Descripción",
        "Especificaciones técnicas",
        "SLA",
        "Cuándo ofrecerlo",
        "Cuándo NO ofrecerlo",
    ]


def test_documento_pmi_a_docx_tiene_portada_y_secciones(tmp_path):
    markdown = (
        "# Acta de Inicio\n\n"
        "## Nombre del proyecto\n"
        "Migración a la nube\n\n"
        "## Patrocinador (sponsor)\n"
        "[Completar]\n"
    )
    buffer = documento_pmi_a_docx("Acta de Inicio", "Cliente Demo S.A.", markdown)

    ruta = tmp_path / "acta.docx"
    ruta.write_bytes(buffer.read())

    documento = Document(str(ruta))
    textos = [p.text for p in documento.paragraphs]

    assert "Acta de Inicio" in textos
    assert "Cliente Demo S.A." in textos
    assert "Nombre del proyecto" in textos
    assert "Migración a la nube" in textos
    # el título no se repite dentro del cuerpo (ya está en la portada)
    assert textos.count("Acta de Inicio") == 1


def test_markdown_a_docx_genera_documento_valido(tmp_path):
    markdown = "# Título\n\n## Sección\n- item uno\n- **item en negrita**\n\ntexto normal"
    buffer = markdown_a_docx(markdown)

    ruta = tmp_path / "salida.docx"
    ruta.write_bytes(buffer.read())

    documento = Document(str(ruta))
    textos = [p.text for p in documento.paragraphs]
    assert "Título" in textos
    assert "Sección" in textos
    assert "item uno" in textos
